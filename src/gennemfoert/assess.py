"""Draft brandteknisk vurdering (DOCX) per group of identical deviations.

The draft carries everything that is fact: the project, the penetrations covered, the tested configuration it is
compared with (ETA section and page), each parameter as designed vs as tested, the tested alternative and the
conditions. The judgement itself is left to the adviser, with the questions it has to answer.
"""
import datetime
import json
from collections import OrderedDict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from . import text

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "out" / "vurderinger"

GUIDE = {
    "diameter": "Er den projekterede dimension inden for det, som de testede dimensioner og tætningens virkemåde kan dække? "
                "Findes en testet løsning med en anden tætningstype (fx brandmanchet) for denne dimension?",
    "godstykkelse": "Hvordan påvirker den afvigende godstykkelse rørets opførsel ved brand (lukning af plastrør, varmeledning i "
                    "metalrør)? Er afvigelsen inden for det, standardens direkte anvendelsesområde (EN 1366-3) tillader?",
    "isoleringstykkelse": "Påvirker den tykkere/tyndere isolering I-kriteriet (temperatur på den ueksponerede side) og tætningens funktion?",
    "isoleringstype": "Kan den projekterede isoleringstype erstattes af den testede, eller skal afvigelsen vurderes?",
    "isoleringsfabrikat": "Er det projekterede isoleringsfabrikat dokumenteret ligeværdigt med det testede (materiale, densitet, "
                          "brandklasse efter EN 13501-1)? Ellers: skift til det testede fabrikat.",
    "klassifikation": "Den testede klassifikation er lavere end kravet. Kan kravet opfyldes med en anden opbygning (fx 2 lag plade)?",
    "rørende": "Svarer rørendebetingelsen (U/U, C/U, U/C, C/C) til systemets art (åbent/lukket)?",
    "afstand mellem tætninger": "Kan tætningerne samles i én tætning, eller flyttes, så afstanden bliver mindst 200 mm?",
    "hulstørrelse": "Kan hullet reduceres til den testede maksimale størrelse?",
    "kantafstand": "Kan installationen placeres, så afstanden til hullets kant overholdes?",
}


def groups(pens):
    g = OrderedDict()
    for p in pens:
        v = p["verdict"]
        if v["status"] not in ("AFVIGELSE", "IKKE_DAEKKET"):
            continue
        h = p["host"]
        key = (text.service_line(p["service"]), h["kind"], h["thickness_mm"], h["rating"], v["status"],
               tuple((v["nearest"] or [{}])[0].get("deviates", [])))
        g.setdefault(key, []).append(p)
    return list(g.values())


def _shade(cell, hex_):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_)
    tc.append(sh)


def draft(group, n, project="Blok C, etageboliger (demonstration)"):
    p0 = group[0]
    s, h, v = p0["service"], p0["host"], p0["verdict"]
    nb = v["nearest"][0] if v["nearest"] else None
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    t = doc.add_heading("Brandteknisk vurdering", 0)
    para = doc.add_paragraph()
    r = para.add_run("UDKAST, udarbejdet af Gennemført. Skal gennemgås, færdiggøres og underskrives af rådgiver.")
    r.bold, r.font.color.rgb = True, RGBColor(0xC2, 0x41, 0x0C)
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Light List"
    for k, val in [("Projekt", project), ("Dato", datetime.date.today().isoformat()), ("Sagsnr.", f"GV-{n:02d}"),
                   ("Gennemføringer", ", ".join(p["id"] for p in group) + f" ({len(group)} stk.)"),
                   ("Gennemføring", text.service_line(s)), ("Konstruktion", f"{h['name']}: {text.HOST[h['kind']]} {h['thickness_mm']} mm"),
                   ("Krav", f"{h['rating']} (tætning EI {h['ei_required']})"), ("System", s.get("system") or "ukendt")]:
        row = meta.add_row().cells
        row[0].text, row[1].text = k, str(val)
    doc.add_heading("1. Formål", 1)
    doc.add_paragraph(f"At vurdere, om gennemføringen kan udføres, så bygningsdelens brandmodstand ({h['rating']}) genetableres, "
                      f"selv om den projekterede udførelse ikke svarer præcist til en testet og klassificeret opbygning.")
    doc.add_heading("2. Grundlag", 1)
    for g in ["ETA-21/0110, Mulcol Multimastic C System (ETA-Danmark A/S, udstedt 2021-01-01), Annex A"
              + (f", afsnit {nb['section']}, side {nb['page']}" if nb else ""),
              "EAD 350454-00-1104 (gennemføringstætninger); EN 1366-3 (prøvning); EN 13501-2 (klassifikation)",
              "Bygningsreglementet (BR18) med vejledning om brandforhold",
              "Projektets IFC-model (konstruktion, installation, isolering, system og udsparing)"]:
        doc.add_paragraph(g, style="List Bullet")
    doc.add_heading("3. Projekteret udførelse sammenholdt med testet opbygning", 1)
    if nb:
        doc.add_paragraph(f"Nærmeste testede opbygning: {nb['eta']} {nb['section']} (side {nb['page']}): {nb['service']}, "
                          f"{nb['insulation']}, {nb['classification']}" + (f", {nb['board']}" if nb.get("board") else "") + ".")
        tb = doc.add_table(rows=1, cols=4)
        tb.style = "Light Grid Accent 1"
        for c, hd in zip(tb.rows[0].cells, ["Parameter", "Projekteret", "Testet", "Overensstemmelse"]):
            c.text = hd
        for c in nb["checks"]:
            row = tb.add_row().cells
            row[0].text, row[1].text, row[2].text = c["check"], str(c["actual"]), str(c["tested"])
            row[3].text = "Ja" if c["ok"] else "AFVIGER"
            _shade(row[3], "D1FAE5" if c["ok"] else "FDE68A")
            if c.get("note"):
                row[2].text += f" ({c['note']})"
    else:
        doc.add_paragraph("Der findes ingen testet opbygning for denne installation i denne konstruktionstype i ETA-21/0110.")
    doc.add_heading("4. Afvigelser", 1)
    devs = [c for c in (nb["checks"] if nb else []) if not c["ok"]]
    for c in devs:
        doc.add_paragraph(f"{c['check'].capitalize()}: projekteret {c['actual']}, testet {c['tested']}.", style="List Bullet")
    if not devs:
        doc.add_paragraph("Installationstypen er ikke omfattet af ETA'ens testede opbygninger i denne konstruktion.", style="List Bullet")
    if v.get("alternatives"):
        a = v["alternatives"][0]
        doc.add_heading("5. Testet alternativ", 1)
        doc.add_paragraph(f"{a['eta']} {a['section']} (side {a['page']}): {a['service']}, {a['insulation']}, {a['classification']}"
                          + (f", {a['board']}" if a.get("board") else "") + ". Kræver ændring af: " + ", ".join(a["deviates"]) + ".")
    doc.add_heading("6. Rådgiverens vurdering", 1)
    para = doc.add_paragraph()
    r = para.add_run("[Udfyldes af rådgiver.] Spørgsmål vurderingen skal besvare:")
    r.italic = True
    for c in devs or [{"check": "diameter"}]:
        if c["check"] in GUIDE:
            doc.add_paragraph(GUIDE[c["check"]], style="List Number")
    for _ in range(4):
        doc.add_paragraph("_" * 95)
    doc.add_heading("7. Forudsætninger og begrænsninger", 1)
    for c in v["conditions"] + ["Vurderingen gælder kun de nævnte gennemføringer med de angivne materialer og dimensioner.",
                                "Udførelse efter producentens monteringsvejledning og dokumentation med foto pr. gennemføring.",
                                text.DISCLAIMER]:
        doc.add_paragraph(c, style="List Bullet")
    doc.add_heading("8. Underskrift", 1)
    doc.add_paragraph("Udarbejdet af: ______________________   Dato: __________\n\nKontrolleret af: ____________________   Dato: __________")
    OUT.mkdir(parents=True, exist_ok=True)
    name = f"GV-{n:02d}_{group[0]['id']}" + (f"_m.fl" if len(group) > 1 else "") + ".docx"
    path = OUT / name
    doc.save(path)
    return path


if __name__ == "__main__":
    reg = json.loads((ROOT / "results" / "register.json").read_text(encoding="utf-8"))
    gs = groups(reg["penetrations"])
    paths = [draft(g, i + 1) for i, g in enumerate(gs)]
    summary = [{"draft": p.name, "ids": [x["id"] for x in g]} for p, g in zip(paths, gs)]
    (ROOT / "results" / "assessments.json").write_text(json.dumps(summary, indent=1, ensure_ascii=False), encoding="utf-8")
    print(len(paths), "drafts covering", sum(len(g) for g in gs), "penetrations")
