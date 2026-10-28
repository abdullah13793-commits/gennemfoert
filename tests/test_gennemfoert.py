"""v1 acceptance tests (SPEC.md A1-A6), run in board-only mode (ETA-21/0110), which reproduces the v1 register exactly.
The full product range and browser parity are tested in test_v2.py. Run from the product root: python -m pytest tests"""
import functools
import http.server
import json
import os
import random
import sys
import threading
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from gennemfoert import engine, pipeline  # noqa: E402

RECS = json.loads((ROOT / "data" / "eta" / "records.json").read_text(encoding="utf-8"))
WEB = ROOT / "web" if (ROOT / "web").exists() else ROOT / "docs" / "v2"          # product folder vs published repo
OUT = ROOT / "out" if (ROOT / "out").exists() else WEB / "downloads"


@pytest.fixture(autouse=True, scope="module")
def board_only():
    engine.use_systems(["Multimastic C"])
    yield
    engine.use_systems()


def Q(host, t, ei, kind="pipe", mat=None, od=None, wall=None, open_=None, ins=None, ap=None):
    q = {"host_kind": host, "thickness_mm": t, "ei_required": ei, "service_kind": kind, "material": mat, "od_mm": od, "wall_mm": wall,
         "open_system": open_, "insulation": engine.norm_insulation(ins), "hollow_core": False}
    if ap:
        q["aperture_mm"] = ap
    return q


SW = lambda t: {"category": "stenuld", "thickness_mm": t, "product": "Stenuld rørskål"}
KF = lambda t, p="Kaiflex ST": {"category": "elastomer", "thickness_mm": t, "product": p}

# Hand-derived from ETA-21/0110 Annex A (page numbers in comments)
GOLDEN = [
    ("G1 PVC 110x3.4 in floor", Q("rigid_floor", 200, 60, mat="pvc-u", od=110, wall=3.4, open_=True), "TESTET", {"section": "A.4.5"}),   # p52
    ("G2 PVC 125 in floor", Q("rigid_floor", 200, 60, mat="pvc-u", od=125, wall=3.9, open_=True), "AFVIGELSE", {"deviates": ["diameter"]}),
    ("G3 Cu 28 Kaiflex 13 light wall", Q("flexible_wall", 100, 60, mat="copper", od=28, wall=1.0, open_=False, ins=KF(13)), "TESTET", {"section": "A.1.6.2"}),  # p20
    # two rows one step away (p20): 9-13 mm gives EI 90, 13-25 mm gives EI 60; both must be named
    ("G4 Cu 28 Kaiflex 20 EI90", Q("flexible_wall", 100, 90, mat="copper", od=28, wall=1.0, open_=False, ins=KF(20)), "AFVIGELSE",
     {"named": ["isoleringstykkelse", "klassifikation"]}),
    ("G5 steel 114.3 SW30 EI120", Q("rigid_wall", 150, 120, mat="steel", od=114.3, wall=3.6, open_=False, ins=SW(30)), "TESTET", {"has": "A.3.4"}),  # p33 + graph p34
    ("G6 PVC 32 75 mm wall EI30", Q("flexible_wall", 75, 30, mat="pvc-u", od=32, wall=1.8, open_=True), "TESTET", {"section": "A.2.1"}),  # p25
    ("G7 PVC 32 75 mm wall EI60", Q("flexible_wall", 75, 60, mat="pvc-u", od=32, wall=1.8, open_=True), "AFVIGELSE", {"deviates": ["klassifikation"]}),
    ("G8 cable 25 mm", Q("rigid_wall", 150, 60, kind="cable", od=25), "TESTET", {}),
    ("G9 steel drain open system", Q("rigid_wall", 150, 60, mat="steel", od=60.3, wall=2.9, open_=True, ins=SW(30)), "AFVIGELSE", {"deviates": ["rørende"]}),
    ("G10 Cu 54 SW20 EI90", Q("rigid_wall", 150, 90, mat="copper", od=54, wall=1.2, open_=False, ins=SW(20)), "TESTET", {"has": "A.3.6"}),  # p36
    ("G11 Cu 54 SW20 EI120 big hole", Q("rigid_wall", 150, 120, mat="copper", od=54, wall=1.2, open_=False, ins=SW(20), ap=[200, 200]),
     "AFVIGELSE", {"deviates": ["hulstørrelse"]}),
    ("G12 Cu 54 SW20 EI120", Q("rigid_wall", 150, 120, mat="copper", od=54, wall=1.2, open_=False, ins=SW(20)), "TESTET", {"section": "A.3.6"}),
    ("G13 PP 75 in floor", Q("rigid_floor", 150, 60, mat="pp", od=75, wall=2.6, open_=True), "IKKE_DAEKKET", {}),
    ("G14 duct", Q("flexible_wall", 100, 60, kind="duct", od=160), "ANDET_SYSTEM", {}),
    ("G15 Cu 22 SW20 in floor", Q("rigid_floor", 150, 60, mat="copper", od=22, wall=1.0, open_=False, ins=SW(20)), "AFVIGELSE",
     {"deviates": ["klassifikation"], "alt": "A.4.8"}),
    ("G16 steel 42.4 SW30 in floor", Q("rigid_floor", 150, 60, mat="steel", od=42.4, wall=2.6, open_=False, ins=SW(30)), "TESTET", {"has": "A.4.4.1"}),
    ("G17 PP 50 light wall", Q("flexible_wall", 100, 60, mat="pp", od=50, wall=1.8, open_=True), "AFVIGELSE", {"deviates": ["diameter"]}),
    ("G18 unknown material", Q("flexible_wall", 100, 60, mat=None, od=22, wall=None, open_=False), "MANGLER_DATA", {}),
]


def test_a1_extraction():
    assert len(RECS) >= 280
    assert len({r["section"] for r in RECS}) >= 36
    assert all(r["page"] >= 8 and r["section"].startswith("A.") for r in RECS)
    assert sum(1 for r in RECS if r["row_bbox"]) / len(RECS) > 0.97
    assert {r["source"] for r in RECS} == {"pdf", "override"}
    # spot checks against the source pages
    p52 = [r for r in RECS if r["page"] == 52]
    assert p52[0]["service"]["od_max"] == 110 and p52[0]["service"]["wall_min"] == 3.4 and p52[0]["classification"]["EI"] == 90
    p20 = [r for r in RECS if r["page"] == 20 and r["service"].get("material") == "copper"]
    assert [r["classification"]["EI"] for r in p20] == [120, 90, 60]


@pytest.mark.parametrize("name,q,status,exp", GOLDEN, ids=[g[0] for g in GOLDEN])
def test_a2_golden(name, q, status, exp):
    v = engine.evaluate(q)
    assert v["status"] == status, (name, v["status"], v["nearest"][:1])
    if "section" in exp:
        assert v["matches"][0]["section"] == exp["section"]
    if "has" in exp:
        assert exp["has"] in {m["section"] for m in v["matches"]}
    if "deviates" in exp:
        assert set(exp["deviates"]) <= set(v["nearest"][0]["deviates"])
    if "named" in exp:
        for d in exp["named"]:
            assert any(n["deviates"] == [d] for n in v["nearest"]), (d, [n["deviates"] for n in v["nearest"]])
    if "alt" in exp:
        assert exp["alt"] in {a["section"] for a in v["alternatives"]}


@pytest.fixture(scope="module")
def run():
    return pipeline.run()


def test_a3_a4_answer_key(run):
    pens, stats = run
    planted = json.loads((ROOT / "models" / "planted.json").read_text(encoding="utf-8"))
    ev = pipeline.evaluate_against(pens, planted)
    assert ev["extra"] == 0 and ev["all_ok"], [r for r in ev["rows"] if not r["ok"]]


def test_a5_writeback():
    wb = json.loads((ROOT / "results" / "writeback.json").read_text())
    assert wb["seals"] == wb["with_psets"] == wb["fills_opening"] == 58
    with zipfile.ZipFile(OUT / "Gennemfoert_afvigelser.bcfzip") as z:
        assert sum(n.endswith("markup.bcf") for n in z.namelist()) == wb["bcf_topics"] >= 19     # 22 board-only, 19 full range


def test_a6_register():
    from openpyxl import load_workbook
    ws = load_workbook(OUT / "Gennemfoeringsregister.xlsx")["Register"]
    assert ws.max_row - 1 == 58
    assert (OUT / "Gennemfoeringsregister.pdf").stat().st_size > 50_000
